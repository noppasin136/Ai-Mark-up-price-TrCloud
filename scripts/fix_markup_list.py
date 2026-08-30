"""Apply confirmed Category / Markup corrections to markup_list.xlsx.

These are one-off fixes signed off by finance (see docs/COSTING_MODEL.md §4.3):

  702-0026, 602-0010, 704-0014  Category  WH -> Freeze
      Frozen items mis-tagged as warehouse stock. They already carry the
      Freeze rate (0.20); only the Category label was wrong.
  502-0007                      Markup    0.12 -> 0.25
      Central Kitchen item that should be marked up like the rest of the
      Central Kitchen catalogue.

Only markup_list.xlsx is touched — Category and Markup do not exist in
sale_list.xlsx. The original is copied to data/input/archive/ first and a
change log is written to data/output/.

    python scripts/fix_markup_list.py --dry-run     # show the changes only
    python scripts/fix_markup_list.py               # apply them

Safe to re-run: a correction already applied is reported and skipped.
"""

from __future__ import annotations

import argparse
import datetime as dt
import shutil
import sys
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
IN = ROOT / "data" / "input"
ARCHIVE = IN / "archive"
TARGET = "markup_list.xlsx"
SHEET = "SellingPrice"
CODE_HEADER = "Code"

# sku, column, expected current value, corrected value, why
CORRECTIONS: list[tuple[str, str, object, object, str]] = [
    ("702-0026", "Category", "WH", "Freeze", "frozen item mis-tagged as WH"),
    ("602-0010", "Category", "WH", "Freeze", "frozen item mis-tagged as WH"),
    ("704-0014", "Category", "WH", "Freeze", "frozen item mis-tagged as WH"),
    ("502-0007", "Markup", 0.12, 0.25, "Central Kitchen item should be 25%"),
]


def _same(a: object, b: object) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) < 1e-9
    return str(a).strip() == str(b).strip()


def _locate(ws) -> dict[str, int]:
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    for need in (CODE_HEADER, "Category", "Markup"):
        if need not in headers:
            sys.exit(f"'{need}' column not found in {ws.title}. Headers: {list(headers)}")
    return headers


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    path = IN / TARGET
    if not path.exists():
        sys.exit(f"{TARGET} not found in {IN}")

    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    cols = _locate(ws)
    row_of = {
        str(ws.cell(row=r, column=cols[CODE_HEADER]).value).strip(): r
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=cols[CODE_HEADER]).value is not None
    }

    applied, skipped, unexpected = [], [], []
    for sku, column, expect, new, reason in CORRECTIONS:
        row = row_of.get(sku)
        if row is None:
            unexpected.append((sku, column, "SKU not in markup_list", "", reason))
            continue
        cell = ws.cell(row=row, column=cols[column])
        current = cell.value
        if _same(current, new):
            skipped.append((sku, column, current, new, "already applied"))
            continue
        if not _same(current, expect):
            unexpected.append((sku, column, current, new, f"expected {expect!r}, left alone"))
            continue
        applied.append((sku, column, current, new, reason))
        if not args.dry_run:
            cell.value = new

    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    if applied and not args.dry_run:
        ARCHIVE.mkdir(parents=True, exist_ok=True)
        backup = ARCHIVE / f"{path.stem}_before_category_fix_{stamp}{path.suffix}"
        shutil.copy2(path, backup)
        wb.save(path)
        print(f"{TARGET}: {len(applied)} correction(s) applied  (original -> archive/{backup.name})")
    elif applied:
        print(f"{TARGET}: {len(applied)} correction(s) would be applied (dry run)")
    else:
        print(f"{TARGET}: nothing to change")
    wb.close()

    def _show(title: str, rows: list) -> None:
        if not rows:
            return
        print(f"\n{title}")
        print(
            pd.DataFrame(rows, columns=["sku", "column", "was", "now", "note"]).to_string(index=False)
        )

    _show("Applied:", applied)
    _show("Skipped:", skipped)
    _show("LEFT ALONE — value was not what we expected:", unexpected)

    if applied and not args.dry_run:
        log = ROOT / "data" / "output" / f"markup_list_fixes_{stamp}.csv"
        log.parent.mkdir(parents=True, exist_ok=True)
        pd.DataFrame(applied, columns=["sku", "column", "was", "now", "reason"]).to_csv(
            log, index=False
        )
        print(f"\nChange log: {log}")


if __name__ == "__main__":
    main()
