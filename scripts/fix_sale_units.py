"""Correct the Sale Unit column in markup_list.xlsx and sale_list.xlsx against W10.

Some rows carry a selling unit that does not exist in W10 for that SKU (most often
"bag" where W10 knows only "row", "Kg.", "stick" and so on). Where W10 lists exactly
ONE unit for the SKU, the correct value is unambiguous and is written in.

Rows where W10 offers a choice are left untouched and reported — those need a human.
Originals are copied to data/input/archive/ before anything is written.

    python scripts/fix_sale_units.py --dry-run     # show the changes only
    python scripts/fix_sale_units.py               # apply them
"""

from __future__ import annotations

import argparse
import datetime as dt
import shutil
import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.worksheet import _reader

# The ERP writes the text "null" into some numeric cells; recover as blank.
_orig_cast = _reader._cast_number


def _safe_cast(value):
    try:
        return _orig_cast(value)
    except (ValueError, TypeError):
        return None


_reader._cast_number = _safe_cast

ROOT = Path(__file__).resolve().parent.parent
IN = ROOT / "data" / "input"
ARCHIVE = IN / "archive"
TARGETS = ("markup_list.xlsx", "sale_list.xlsx")
SHEET = "SellingPrice"
SKU_HEADER, UNIT_HEADER = "Code", "Sale Unit"


def w10_units() -> dict[str, list[tuple[str, float, float, bool]]]:
    """sku -> [(unit, coefficient, price, is_base), ...]"""
    src = sorted(IN.glob("W10*.xlsx"))
    if not src:
        sys.exit("No W10 export found in data/input/")
    w = pd.read_excel(src[0], sheet_name="data")
    w.columns = [str(c).strip() for c in w.columns]
    w = w.rename(columns={"รหัสสินค้า": "sku", "หน่วย": "uom", "Ccoefficient": "coef",
                          "ราคาขาย": "price", "เป็นหน่วยหลัก": "is_main"})
    w["sku"] = w["sku"].astype(str).str.strip()
    w["uom"] = w["uom"].astype(str).str.strip()
    for c in ("coef", "price", "is_main"):
        w[c] = pd.to_numeric(w[c], errors="coerce")
    out: dict[str, list] = {}
    for r in w.itertuples():
        out.setdefault(r.sku, []).append((r.uom, r.coef, r.price, r.is_main == 1))
    return out


def locate(ws) -> tuple[int, int]:
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    for need in (SKU_HEADER, UNIT_HEADER):
        if need not in headers:
            sys.exit(f"'{need}' column not found in {ws.title}. Headers: {list(headers)}")
    return headers[SKU_HEADER], headers[UNIT_HEADER]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    units = w10_units()
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    changes, ambiguous, unknown = [], [], []

    for filename in TARGETS:
        path = IN / filename
        if not path.exists():
            print(f"skip {filename} (not present)")
            continue

        wb = openpyxl.load_workbook(path)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        sku_col, unit_col = locate(ws)
        edits = 0

        for row in range(2, ws.max_row + 1):
            sku = ws.cell(row=row, column=sku_col).value
            if sku is None:
                continue
            sku = str(sku).strip()
            cell = ws.cell(row=row, column=unit_col)
            current = "" if cell.value is None else str(cell.value).strip()

            options = units.get(sku)
            if not options:
                unknown.append((filename, sku, current))
                continue
            if current and current.lower() in {u.lower() for u, *_ in options}:
                continue  # already valid

            if len(options) == 1:
                correct, coef, price, _ = options[0]
                changes.append(
                    {"file": filename, "row": row, "sku": sku,
                     "was": current or "(blank)", "now": correct,
                     "w10_price": price, "coefficient": coef}
                )
                if not args.dry_run:
                    cell.value = correct
                edits += 1
            else:
                ambiguous.append(
                    {"file": filename, "sku": sku, "was": current or "(blank)",
                     "options": ", ".join(f"{u}(c={c:g})" for u, c, *_ in options)}
                )

        if edits and not args.dry_run:
            ARCHIVE.mkdir(parents=True, exist_ok=True)
            backup = ARCHIVE / f"{path.stem}_before_unit_fix_{stamp}{path.suffix}"
            shutil.copy2(path, backup)
            wb.save(path)
            print(f"{filename}: {edits} unit(s) corrected  (original -> archive/{backup.name})")
        elif edits:
            print(f"{filename}: {edits} unit(s) would be corrected (dry run)")
        else:
            print(f"{filename}: nothing to change")
        wb.close()

    if changes:
        df = pd.DataFrame(changes)
        print("\nCorrections:")
        print(df.drop_duplicates(subset=["sku", "was", "now"])[
            ["sku", "was", "now", "w10_price", "coefficient"]].to_string(index=False))
        if not args.dry_run:
            log = ROOT / "data" / "output" / f"sale_unit_fixes_{stamp}.csv"
            log.parent.mkdir(parents=True, exist_ok=True)
            df.to_csv(log, index=False)
            print(f"\nChange log: {log}")

    if ambiguous:
        print("\nLEFT ALONE — W10 offers more than one unit, so the correct one is a judgement call:")
        print(pd.DataFrame(ambiguous).drop_duplicates(subset=["sku"]).to_string(index=False))
    if unknown:
        print(f"\nNot found in W10 at all: {len(set(s for _, s, _ in unknown))} SKU(s)")


if __name__ == "__main__":
    main()
