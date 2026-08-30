"""Price review — an opt-out gate for soft-flagged rows.

Guardrail breaches and other soft flags (docs/COSTING_MODEL.md §8.C) ship on the
Price Upload by default. This sheet lets a person veto a specific row: set its
``Decision`` to ``HOLD`` and that SKU is kept off the next upload. Silence — a
blank or ``OK`` — means the price ships.

``markup update`` refreshes the sheet with the current soft-flagged rows and
never overwrites a decision already made, exactly like ``unit_review.xlsx``.
"""

from __future__ import annotations

import datetime as dt
import logging
from pathlib import Path

import pandas as pd

log = logging.getLogger(__name__)

REVIEW_FILENAME = "price_review.xlsx"
SHEET = "Price Review"
HOLD, OK = "HOLD", "OK"

COLUMNS = [
    "SKU", "Product Name", "Category", "Cost Source", "Current Price",
    "Suggested Price", "Change %", "Flags", "Decision", "Note", "First Seen",
]


def review_path(cfg) -> Path:
    return cfg.resolve(Path("config") / REVIEW_FILENAME)


def load_holds(cfg) -> set[str]:
    """SKUs whose Decision is HOLD — kept off the upload. Missing file = none."""
    path = review_path(cfg)
    if not path.exists():
        return set()
    df = pd.read_excel(path, sheet_name=SHEET)
    df.columns = [str(c).strip() for c in df.columns]
    if "SKU" not in df.columns or "Decision" not in df.columns:
        return set()
    held = df[df["Decision"].astype(str).str.strip().str.upper() == HOLD]
    holds = {str(s).strip() for s in held["SKU"] if str(s).strip().lower() not in ("", "nan")}
    if holds:
        log.info("Price review: %d SKU(s) held back by a HOLD decision", len(holds))
    return holds


def sync(cfg, rows: pd.DataFrame) -> tuple[Path, int, int]:
    """Merge the current soft-flagged rows into the sheet, keeping decisions.

    ``rows`` needs: sku, product_name, category, cost_source, current_price,
    suggested_price, change_pct, flag_codes. Returns ``(path, added, held)``.
    """
    path = review_path(cfg)
    today = dt.date.today().isoformat()

    existing = pd.DataFrame(columns=COLUMNS)
    if path.exists():
        existing = pd.read_excel(path, sheet_name=SHEET)
        existing.columns = [str(c).strip() for c in existing.columns]
    decided = {
        str(r.get("SKU")).strip(): (str(r.get("Decision") or "").strip(), str(r.get("Note") or "").strip())
        for r in existing.to_dict("records")
        if str(r.get("SKU") or "").strip().lower() not in ("", "nan")
    }
    first_seen = {
        str(r.get("SKU")).strip(): r.get("First Seen")
        for r in existing.to_dict("records")
    }

    out = []
    for r in rows.itertuples():
        prior_decision, prior_note = decided.get(str(r.sku), ("", ""))
        out.append({
            "SKU": r.sku,
            "Product Name": getattr(r, "product_name", None),
            "Category": getattr(r, "category", None),
            "Cost Source": getattr(r, "cost_source", None),
            "Current Price": getattr(r, "current_price", None),
            "Suggested Price": getattr(r, "suggested_price", None),
            "Change %": getattr(r, "change_pct", None),
            "Flags": getattr(r, "flag_codes", None),
            "Decision": prior_decision or None,
            "Note": prior_note or None,
            "First Seen": first_seen.get(str(r.sku)) or today,
        })

    merged = pd.DataFrame(out, columns=COLUMNS) if out else pd.DataFrame(columns=COLUMNS)
    merged = merged.sort_values("SKU").reset_index(drop=True)
    _write(path, merged)

    held = int(merged["Decision"].astype(str).str.strip().str.upper().eq(HOLD).sum())
    added = sum(1 for r in out if str(r["SKU"]) not in decided)
    return path, added, held


def _write(path: Path, df: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        df.to_excel(xl, sheet_name=SHEET, index=False)
        ws = xl.book[SHEET]

        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

        widths = {"SKU": 12, "Product Name": 36, "Category": 14, "Cost Source": 13,
                  "Current Price": 13, "Suggested Price": 14, "Change %": 10,
                  "Flags": 46, "Decision": 12, "Note": 32, "First Seen": 12}
        for i, col in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)

        dcol = get_column_letter(list(df.columns).index("Decision") + 1)
        dv = DataValidation(
            type="list", formula1=f'"{OK},{HOLD}"', allow_blank=True, showDropDown=False
        )
        dv.error = "Choose OK or HOLD (blank means the price ships)"
        ws.add_data_validation(dv)
        dv.add(f"{dcol}2:{dcol}{max(len(df) + 1, 2)}")

        flags_idx = list(df.columns).index("Flags") + 1
        for r in range(2, len(df) + 2):
            ws.cell(row=r, column=flags_idx).alignment = Alignment(wrap_text=True, vertical="top")
            key = str(ws.cell(row=r, column=list(df.columns).index("Decision") + 1).value or "").strip().upper()
            if key == HOLD:
                for c in range(1, len(df.columns) + 1):
                    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FCE4D6")
