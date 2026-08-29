"""Excel writer — the final deliverable.

Sheet layout:
  Price Upload  lean, ERP-import ready: SKU | Unit | Sale Price
  Detail        every column behind each price
  Exceptions    rows a human must clear before upload
  Cost Audit    the GR2 lines each cost was built from
  Run Summary   the exact parameters this file was produced with
"""

from __future__ import annotations

import datetime as dt
import logging
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import AppConfig
from ..pipeline import RunResult

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
FLAG_FILL = PatternFill("solid", fgColor="FCE4E4")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY_COLS = {
    "unit_cost", "raw_price", "suggested_price", "parallel_price", "current_price",
    "price_change", "min_line_cost", "max_line_cost", "effective_cost", "Sale Price",
    "freight", "duty", "other_landed",
    "unit_cost_prev", "suggested_price_prev", "cost_delta", "price_delta",
    "avg_cost", "avg_current_price", "avg_new_price",
}
PCT_COLS = {
    "markup_pct", "margin_pct", "change_pct", "min_margin_pct",
    "price_delta_pct", "markup_pct_prev", "avg_markup_pct", "avg_margin_pct",
    "avg_change_pct",
}

REPORT_TITLES = {
    "summary": "Summary",
    "movers": "Biggest Movers",
    "category": "By Category",
    "comparison": "vs Previous Run",
    "comparison_summary": "Comparison Summary",
}

SHEET_TITLES = {
    "price_upload": "Price Upload",
    "detail": "Detail",
    "exceptions": "Exceptions",
    "cost_audit": "Cost Audit",
    "run_summary": "Run Summary",
}


def _resolve_path(cfg: AppConfig) -> Path:
    name = cfg.out_filename.format(
        method=cfg.method,
        period=cfg.period_days,
        timestamp=dt.datetime.now().strftime("%Y%m%d_%H%M%S"),
        as_of=cfg.as_of_date,
    )
    out_dir = cfg.resolve(cfg.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / name


def write_workbook(result: RunResult, cfg: AppConfig, path: str | Path | None = None) -> Path:
    path = Path(path) if path else _resolve_path(cfg)

    frames = {
        "price_upload": result.upload,
        "detail": result.detail,
        "exceptions": result.exceptions,
        "cost_audit": result.cost_audit,
        "run_summary": result.summary,
    }
    upload_title = cfg.upload_opts.get("sheet_name", "Price Upload")

    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        for key in cfg.sheets:
            df = frames.get(key)
            if df is None:
                log.warning("Unknown sheet '%s' in output.sheets — skipped", key)
                continue
            title = upload_title if key == "price_upload" else SHEET_TITLES.get(key, key)
            (df if not df.empty else pd.DataFrame({"(no rows)": []})).to_excel(
                writer, sheet_name=title[:31], index=False
            )
            _style(writer.book[title[:31]], df, cfg, key)

    log.info("Workbook written: %s", path)
    return path


def write_report_workbook(views: dict, cfg: AppConfig, path: Path) -> Path:
    """Write the analysis workbook produced by ``markup report``."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        for key, df in views.items():
            if df is None:
                continue
            title = REPORT_TITLES.get(key, key)[:31]
            (df if not df.empty else pd.DataFrame({"(no rows)": []})).to_excel(
                writer, sheet_name=title, index=False
            )
            _style(writer.book[title], df, cfg, "run_summary" if "summary" in key else key)
    log.info("Report written: %s", path)
    return path


def _style(ws, df: pd.DataFrame, cfg: AppConfig, key: str) -> None:
    if df.empty:
        return

    headers = list(df.columns)
    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

        letter = get_column_letter(col_idx)
        width = max(len(str(name)) + 4, min(int(df[name].astype(str).str.len().max() or 10) + 3, 42))
        ws.column_dimensions[letter].width = width

        if name in MONEY_COLS:
            fmt = "#,##0.00"
        elif name in PCT_COLS:
            fmt = "0.00"
        else:
            fmt = None
        if fmt:
            for row_idx in range(2, len(df) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = fmt

    ws.row_dimensions[1].height = 28
    if cfg.freeze_header:
        ws.freeze_panes = "A2"
    if cfg.autofilter:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df) + 1}"

    if key in {"detail", "exceptions"} and "flag_codes" in headers:
        blocked_col = headers.index("blocked") + 1 if "blocked" in headers else None
        flag_col = headers.index("flag_codes") + 1
        for row_idx in range(2, len(df) + 2):
            if not ws.cell(row=row_idx, column=flag_col).value:
                continue
            is_blocked = bool(ws.cell(row=row_idx, column=blocked_col).value) if blocked_col else True
            fill = FLAG_FILL if is_blocked else WARN_FILL
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    if key == "run_summary":
        for row_idx in range(2, len(df) + 2):
            ws.cell(row=row_idx, column=1).font = Font(bold=True)
